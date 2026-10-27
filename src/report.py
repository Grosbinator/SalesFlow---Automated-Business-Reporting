"""Génération du rapport Excel reports/monthly_report.xlsx.

Feuilles :
- Dashboard     : indicateurs clés et graphiques ;
- Revenue       : évolution mensuelle des paiements réussis ;
- Customers     : liste et statistiques des clients ;
- Subscriptions : répartition par formule et par statut ;
- Data Quality  : résumé des anomalies détectées par cleaning.py.

Les chiffres viennent de analytics.py (agrégations SQL) ; ce module ne fait
que la mise en forme (openpyxl) et les graphiques (matplotlib).

Usage : python src/report.py   (après cleaning.py et database.py)
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # pas de fenêtre : les graphiques deviennent des images

import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.ticker import FuncFormatter  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import analytics  # noqa: E402
from cleaning import ANOMALIES_FILE, ANOMALY_COLUMNS  # noqa: E402

ROOT_DIR = Path(__file__).resolve().parent.parent
REPORT_FILE = ROOT_DIR / "reports" / "monthly_report.xlsx"

# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

# Couleurs des graphiques : palette lisible y compris par les daltoniens.
BLUE = "#2a78d6"
ORANGE = "#eb6834"
INK = "#0b0b0b"
INK_MUTED = "#52514e"
GRID = "#e4e3df"

TITLE_FONT = Font(size=16, bold=True, color="1F3A5F")
SUBTITLE_FONT = Font(size=10, italic=True, color="52514E")
SECTION_FONT = Font(size=12, bold=True, color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
ZEBRA_FILL = PatternFill("solid", fgColor="F4F6F9")
TILE_FILL = PatternFill("solid", fgColor="EEF3FA")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

EURO = '#,##0.00 "€"'
PERCENT = "0.0%"
INTEGER = "0"

# Libellés français des colonnes renvoyées par analytics.py.
LABELS = {
    "indicator": "Indicateur",
    "value": "Valeur",
    "month": "Mois",
    "recurring_revenue": "Revenu récurrent théorique",
    "collected_revenue": "Revenu encaissé",
    "gap": "Écart",
    "running_subscriptions": "Abonnements en cours",
    "successful_payments": "Paiements réussis",
    "plan": "Formule",
    "subscriptions": "Abonnements",
    "active_subscriptions": "Abonnements actifs",
    "share_pct": "Part (%)",
    "active": "Actifs",
    "cancelled": "Résiliés",
    "total": "Total",
    "customer_id": "ID client",
    "company_name": "Entreprise",
    "email": "Email",
    "country": "Pays",
    "signup_date": "Inscription",
    "active_plans": "Formules actives",
    "total_paid": "Total payé",
    "last_payment_date": "Dernier paiement",
    "is_active": "Client actif",
    "customers": "Clients",
    "active_customers": "Clients actifs",
    "file_name": "Fichier",
    "row_number": "Ligne",
    "column_name": "Colonne",
    "error_message": "Message",
    "error_type": "Type d'erreur",
    "anomalies": "Anomalies",
}
EURO_COLUMNS = {"recurring_revenue", "collected_revenue", "gap", "total_paid"}


# ---------------------------------------------------------------------------
# Outils d'écriture dans les feuilles
# ---------------------------------------------------------------------------

def to_excel_value(value):
    """Convertit une valeur pandas/numpy en valeur que openpyxl sait écrire."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if hasattr(value, "item"):  # types numpy -> types Python
        value = value.item()
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    return value


def write_title(sheet, title, subtitle=None):
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    if subtitle:
        sheet["A2"] = subtitle
        sheet["A2"].font = SUBTITLE_FONT


def write_table(sheet, df, row, column=1, title=None, formats=None):
    """Écrit un DataFrame sous forme de tableau mis en forme.

    `formats` associe un nom de colonne à un format de nombre Excel (les
    colonnes en euros sont formatées automatiquement). Renvoie la première
    ligne libre sous le tableau.
    """
    formats = {**{name: EURO for name in EURO_COLUMNS}, **(formats or {})}

    if title:
        sheet.cell(row, column, title).font = SECTION_FONT
        row += 1

    for offset, name in enumerate(df.columns):
        cell = sheet.cell(row, column + offset, LABELS.get(name, name))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for line, values in enumerate(df.itertuples(index=False), start=1):
        for offset, (name, value) in enumerate(zip(df.columns, values)):
            cell = sheet.cell(row + line, column + offset, to_excel_value(value))
            cell.border = BORDER
            if name in formats:
                cell.number_format = formats[name]
            if line % 2 == 0:
                cell.fill = ZEBRA_FILL

    if df.empty:
        sheet.cell(row + 1, column, "Aucune donnée").font = SUBTITLE_FONT

    return row + max(len(df), 1) + 2


def write_kpi(sheet, row, column, label, value, number_format):
    """Écrit une tuile KPI sur deux colonnes : libellé puis valeur."""
    for line in (row, row + 1):
        sheet.merge_cells(
            start_row=line, start_column=column,
            end_row=line, end_column=column + 1,
        )
        for offset in (0, 1):
            cell = sheet.cell(line, column + offset)
            cell.fill = TILE_FILL
            cell.border = BORDER

    label_cell = sheet.cell(row, column, label)
    label_cell.font = Font(size=10, color="52514E")
    label_cell.alignment = Alignment(horizontal="center")

    value_cell = sheet.cell(row + 1, column, "n/d" if value is None else value)
    value_cell.font = Font(size=18, bold=True, color="1F3A5F")
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.number_format = number_format
    sheet.row_dimensions[row + 1].height = 32


def autosize_columns(sheet, minimum=10, maximum=45):
    """Ajuste la largeur des colonnes au contenu des tableaux (cellules bordées)."""
    widths = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None or cell.border.left.style is None:
                continue
            length = len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, length in widths.items():
        width = min(max(length + 2, minimum), maximum)
        sheet.column_dimensions[get_column_letter(column)].width = width
