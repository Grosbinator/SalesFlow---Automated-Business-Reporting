"""Tests de src/report.py.

Le classeur est construit à partir d'une base SQLite en mémoire (même jeu
de données que tests/test_analytics.py) et d'anomalies fabriquées en
mémoire : data/, anomalies/ et reports/ ne sont jamais touchés.

Lancer les tests depuis la racine du projet :  python -m pytest -v
"""

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import analytics  # noqa: E402
import report  # noqa: E402
from test_analytics import (  # noqa: E402
    CUSTOMERS,
    PAYMENTS,
    SUBSCRIPTIONS,
    load,
    new_connection,
)

ANOMALIES = pd.DataFrame(
    [
        ("payments.csv", 18, "amount", "Valeur manquante"),
        ("payments.csv", 20, "payment_id", "Identifiant dupliqué : P019"),
        ("payments.csv", 21, "payment_id", "Identifiant dupliqué : P019"),
        ("subscriptions.csv", 17, "customer_id",
         "customer_id inexistant dans customers.csv : C999"),
    ],
    columns=report.ANOMALY_COLUMNS,
)


@pytest.fixture
def connection():
    connection = new_connection()
    load(connection, CUSTOMERS, SUBSCRIPTIONS, PAYMENTS)
    yield connection
    connection.close()


def column_values(sheet, column, first_row, last_row):
    return [sheet.cell(row, column).value for row in range(first_row, last_row + 1)]


def find_table(sheet, title):
    """Renvoie les lignes (valeurs) du tableau placé sous le titre `title`."""
    for row in sheet.iter_rows():
        if row[0].value == title:
            start = row[0].row + 2  # titre, puis en-tête
            break
    else:
        raise AssertionError(f"Tableau introuvable : {title}")
    rows = []
    for values in sheet.iter_rows(min_row=start, values_only=True):
        if values[0] is None:
            break
        rows.append(tuple(v for v in values if v is not None))
    return rows


# ---------------------------------------------------------------------------
# 1. Structure du classeur
# ---------------------------------------------------------------------------

def test_report_has_the_five_sheets(connection):
    workbook = report.build_report(connection, ANOMALIES)
    assert workbook.sheetnames == [
        "Dashboard", "Revenue", "Customers", "Subscriptions", "Data Quality",
    ]


def test_main_writes_the_excel_file(connection, tmp_path, monkeypatch):
    """main() enregistre le classeur (dossier reports/ créé si besoin)."""
    output = tmp_path / "reports" / "monthly_report.xlsx"
    monkeypatch.setattr(report, "REPORT_FILE", output)
    monkeypatch.setattr(report, "ANOMALIES_FILE", tmp_path / "absent.csv")
    monkeypatch.setattr(analytics, "get_connection", lambda: connection)

    report.main()

    assert load_workbook(output).sheetnames[0] == "Dashboard"


# ---------------------------------------------------------------------------
# 2. Contenu des feuilles
# ---------------------------------------------------------------------------

def test_dashboard_shows_kpis_and_charts(connection):
    sheet = report.build_report(connection, ANOMALIES)["Dashboard"]
    # Tuiles KPI sur deux colonnes : valeurs en ligne 5, colonnes A, C, E...
    values = [sheet.cell(5, column).value for column in (1, 3, 5, 7, 9, 11)]
    assert values == [216, 108, 1, pytest.approx(54), pytest.approx(0.2), 4]
    assert len(sheet._images) == 2


def test_revenue_sheet_lists_each_month(connection):
    sheet = report.build_report(connection, ANOMALIES)["Revenue"]
    assert sheet.cell(4, 1).value == "Mois"
    assert column_values(sheet, 1, 5, 7) == ["2026-01", "2026-02", "2026-03"]
    assert column_values(sheet, 3, 5, 7) == [29, 108, 79]  # revenu encaissé
    assert sheet.cell(5, 3).number_format == report.EURO


def test_customers_sheet_lists_every_customer(connection):
    sheet = report.build_report(connection, ANOMALIES)["Customers"]
    customers = find_table(sheet, "Liste des clients")
    assert [row[0] for row in customers] == ["C1", "C2", "C3"]
    assert customers[0][-1] == "Oui"  # C1 est actif
    assert customers[2][-1] == "Non"  # C3 n'a aucun abonnement


def test_subscriptions_sheet_crosses_plan_and_status(connection):
    sheet = report.build_report(connection, ANOMALIES)["Subscriptions"]
    assert find_table(sheet, "Par formule et par statut") == [
        ("Pro", 1, 1, 2),
        ("Starter", 1, 0, 1),
    ]


def test_data_quality_summarises_anomalies(connection):
    sheet = report.build_report(connection, ANOMALIES)["Data Quality"]
    assert find_table(sheet, "Par fichier") == [
        ("payments.csv", 3),
        ("subscriptions.csv", 1),
    ]
    assert find_table(sheet, "Par type d'erreur")[0] == ("Identifiant dupliqué", 2)
    assert len(find_table(sheet, "Détail des anomalies")) == 4


# ---------------------------------------------------------------------------
# 3. Cas limites
# ---------------------------------------------------------------------------

def test_missing_anomalies_file_is_reported(connection):
    """Sans anomalies.csv, le rapport est quand même généré avec un message."""
    workbook = report.build_report(connection, None)
    assert "cleaning.py" in workbook["Data Quality"]["A4"].value
    assert workbook["Dashboard"].cell(5, 11).value == "n/d"


def test_load_anomalies_returns_none_when_file_is_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(report, "ANOMALIES_FILE", tmp_path / "absent.csv")
    assert report.load_anomalies() is None


def test_empty_database_still_builds_a_report():
    connection = new_connection()
    try:
        workbook = report.build_report(connection, ANOMALIES.iloc[0:0])
    finally:
        connection.close()
    assert workbook["Dashboard"].cell(5, 1).value == 0
    assert workbook["Dashboard"]._images == []


def test_to_excel_value_converts_booleans_and_missing_values():
    assert report.to_excel_value(True) == "Oui"
    assert report.to_excel_value(float("nan")) is None
