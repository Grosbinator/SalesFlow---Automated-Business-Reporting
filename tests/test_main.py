"""Tests de main.py (pipeline complet).

Les tests de bout en bout redirigent tous les chemins (data/, anomalies/,
reports/) vers un dossier temporaire : les vrais fichiers ne sont jamais
touchés.

Lancer les tests depuis la racine du projet :  python -m pytest -v
"""

import sqlite3
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(ROOT_DIR / "src"))

import cleaning  # noqa: E402
import database  # noqa: E402
import import_data  # noqa: E402
import main  # noqa: E402
import report  # noqa: E402
import validation  # noqa: E402


# ---------------------------------------------------------------------------
# 1. Enchaînement des étapes
# ---------------------------------------------------------------------------

def test_steps_run_in_order():
    calls = []
    steps = [(name, lambda name=name: calls.append(name), None) for name in "ABC"]
    assert main.run_pipeline(steps, show=lambda text="": None) == 0
    assert calls == ["A", "B", "C"]


def test_failure_stops_the_pipeline_with_explicit_message():
    calls = []
    output = []

    def broken():
        raise FileNotFoundError("Fichier introuvable : customers.csv")

    steps = [
        ("Import", lambda: calls.append("import"), None),
        ("Chargement", broken, "Vérifier les données."),
        ("Rapport", lambda: calls.append("rapport"), None),
    ]
    assert main.run_pipeline(steps, show=output.append) == 1
    assert calls == ["import"]  # le rapport n'est jamais lancé
    text = "\n".join(output)
    assert "ÉCHEC à l'étape 2/3 (Chargement)" in text
    assert "FileNotFoundError : Fichier introuvable : customers.csv" in text
    assert "Piste : Vérifier les données." in text


def test_skip_import_removes_the_first_step():
    titles = [title for title, _, _ in main.build_steps(skip_import=True)]
    assert titles == [
        "Nettoyage et validation",
        "Chargement dans SQLite",
        "Calcul des indicateurs",
        "Génération du rapport Excel",
    ]
    assert len(main.build_steps()) == 5


# ---------------------------------------------------------------------------
# 2. Bout en bout sur de petits CSV (dossier temporaire)
# ---------------------------------------------------------------------------

CUSTOMERS = (
    "customer_id,company_name,email,signup_date,country\n"
    "C001,NovaTech,contact@novatech.example,2026-01-12,France\n"
    "C002,Greenlytic,hello@greenlytic.example,2026-02-19,France\n"
)
SUBSCRIPTIONS = (
    "subscription_id,customer_id,plan,monthly_price,start_date,end_date,status\n"
    "S001,C001,Starter,29.0,2026-01-12,,active\n"
    "S002,C002,Pro,79.0,2026-02-19,,active\n"
)
PAYMENTS = (
    "payment_id,subscription_id,payment_date,amount,status,payment_method\n"
    "P001,S001,2026-01-12,29.0,success,card\n"
    "P002,S002,2026-02-19,79.0,success,card\n"
    "P003,S999,2026-02-20,49.0,success,card\n"  # abonnement inconnu
)


@pytest.fixture
def project(tmp_path, monkeypatch):
    """Projet complet dans un dossier temporaire, avec trois CSV bruts."""
    raw = tmp_path / "data" / "raw"
    raw.mkdir(parents=True)
    (raw / "customers.csv").write_text(CUSTOMERS, encoding="utf-8")
    (raw / "subscriptions.csv").write_text(SUBSCRIPTIONS, encoding="utf-8")
    (raw / "payments.csv").write_text(PAYMENTS, encoding="utf-8")

    anomalies = tmp_path / "anomalies"
    processed = tmp_path / "data" / "processed"
    monkeypatch.setattr(cleaning, "RAW_DIR", raw)
    monkeypatch.setattr(import_data, "RAW_DIR", raw)
    monkeypatch.delenv(import_data.SOURCE_ENV_VAR, raising=False)
    monkeypatch.setattr(cleaning, "ANOMALIES_DIR", anomalies)
    monkeypatch.setattr(cleaning, "ANOMALIES_FILE", anomalies / "anomalies.csv")
    monkeypatch.setattr(report, "ANOMALIES_FILE", anomalies / "anomalies.csv")
    monkeypatch.setattr(validation, "PROCESSED_DIR", processed)
    monkeypatch.setattr(validation, "CORRECTIONS_FILE", anomalies / "corrections.csv")
    monkeypatch.setattr(database, "PROCESSED_DIR", processed)
    monkeypatch.setattr(database, "DB_FILE", tmp_path / "data" / "salesflow.db")
    monkeypatch.setattr(report, "REPORT_FILE", tmp_path / "reports" / "monthly_report.xlsx")
    return tmp_path


def test_full_pipeline_produces_the_report(project):
    assert main.main(["--auto", "--skip-import"]) == 0

    workbook = load_workbook(project / "reports" / "monthly_report.xlsx")
    assert workbook.sheetnames == [
        "Dashboard", "Revenue", "Customers", "Subscriptions", "Data Quality",
    ]
    # Le paiement lié à un abonnement inconnu a été écarté avant la base.
    connection = sqlite3.connect(project / "data" / "salesflow.db")
    payments = [row[0] for row in connection.execute("SELECT payment_id FROM payments")]
    connection.close()
    assert sorted(payments) == ["P001", "P002"]
    # Mais il reste signalé dans le rapport d'anomalies.
    anomalies = pd.read_csv(project / "anomalies" / "anomalies.csv", encoding="utf-8-sig")
    assert list(anomalies["column_name"]) == ["subscription_id"]


def test_critical_error_produces_no_report(project, capsys):
    (project / "data" / "raw" / "customers.csv").unlink()

    assert main.main(["--auto", "--skip-import"]) == 1

    assert not (project / "reports" / "monthly_report.xlsx").exists()
    output = capsys.readouterr().out
    assert "ÉCHEC à l'étape 1/4 (Nettoyage et validation)" in output
    assert "FileNotFoundError" in output


def test_running_twice_gives_the_same_database(project):
    assert main.main(["--auto", "--skip-import"]) == 0
    assert main.main(["--auto", "--skip-import"]) == 0
    connection = sqlite3.connect(project / "data" / "salesflow.db")
    assert connection.execute("SELECT COUNT(*) FROM payments").fetchone() == (2,)
    connection.close()


# ---------------------------------------------------------------------------
# 3. Dossier source (--source)
# ---------------------------------------------------------------------------

def test_source_option_imports_from_the_given_folder(project):
    """--source : les CSV sont importés depuis ce dossier vers data/raw/."""
    exports = project / "exports"
    exports.mkdir()
    for path in (project / "data" / "raw").iterdir():
        path.replace(exports / path.name)  # data/raw/ vide avant l'import

    assert main.main(["--auto", "--source", str(exports)]) == 0

    assert (project / "data" / "raw" / "payments.csv").is_file()
    assert (project / "reports" / "monthly_report.xlsx").is_file()


def test_missing_source_folder_stops_at_import(project, capsys):
    assert main.main(["--auto", "--source", str(project / "absent")]) == 1

    output = capsys.readouterr().out
    assert "ÉCHEC à l'étape 1/5 (Import des fichiers sources)" in output
    assert "--source DOSSIER" in output  # la piste explique quoi faire
    assert not (project / "reports" / "monthly_report.xlsx").exists()


def test_source_and_skip_import_are_incompatible():
    with pytest.raises(SystemExit):
        main.main(["--auto", "--source", "exports", "--skip-import"])


# ---------------------------------------------------------------------------
# 4. Mode de validation (fenêtre par défaut, console, auto)
# ---------------------------------------------------------------------------

def test_default_mode_opens_the_validation_window(project, monkeypatch):
    """Sans option, la fenêtre de validation est ouverte avant les calculs."""
    opened = []

    def fake_window(use_journal=True):
        opened.append(use_journal)
        validation.auto_validate()  # comme si l'utilisateur avait continué
        return True

    monkeypatch.setattr(main.validation_gui, "gui_validate", fake_window)
    assert main.main(["--skip-import"]) == 0
    assert opened == [True]  # journal utilisé hors mode démo
    assert (project / "reports" / "monthly_report.xlsx").is_file()


def test_cancelled_window_stops_the_pipeline(project, monkeypatch, capsys):
    monkeypatch.setattr(main.validation_gui, "gui_validate", lambda **kwargs: False)

    assert main.main(["--skip-import"]) == 1

    output = capsys.readouterr().out
    assert "validation annulée" in output
    assert not (project / "reports" / "monthly_report.xlsx").exists()


def test_without_display_falls_back_to_automatic_validation(project, monkeypatch, capsys):
    def no_display(**kwargs):
        raise main.validation_gui.tk.TclError("no display name")

    monkeypatch.setattr(main.validation_gui, "gui_validate", no_display)

    assert main.main(["--skip-import"]) == 0
    assert "validation automatique" in capsys.readouterr().out
    assert (project / "reports" / "monthly_report.xlsx").is_file()


def test_interactive_and_auto_are_exclusive():
    with pytest.raises(SystemExit):
        main.main(["--interactive", "--auto"])


# ---------------------------------------------------------------------------
# 5. Mode démo (--demo)
# ---------------------------------------------------------------------------

@pytest.fixture
def demo_project(project, monkeypatch):
    """Données d'exemple (avec P003 lié à un abonnement inconnu) et un
    journal qui corrige déjà cette erreur."""
    sample = project / "data" / "sample"
    sample.mkdir()
    for path in (project / "data" / "raw").iterdir():
        (sample / path.name).write_bytes(path.read_bytes())
    monkeypatch.setattr(import_data, "SAMPLE_DIR", sample)

    journal = project / "anomalies" / "corrections.csv"
    journal.parent.mkdir()
    pd.DataFrame([{
        "date": "2026-10-28 10:00:00", "file_name": "payments.csv",
        "row_number": 4, "column_name": "subscription_id",
        "error_message": "subscription_id inexistant dans subscriptions.csv : S999",
        "action": "valeur modifiée", "old_value": "S999", "new_value": "S001",
    }], columns=validation.CORRECTION_COLUMNS).to_csv(
        journal, index=False, encoding="utf-8-sig")
    return project


def payment_ids(project):
    connection = sqlite3.connect(project / "data" / "salesflow.db")
    ids = sorted(row[0] for row in connection.execute("SELECT payment_id FROM payments"))
    connection.close()
    return ids


def test_normal_run_applies_the_journal(demo_project):
    assert main.main(["--auto", "--skip-import"]) == 0
    assert payment_ids(demo_project) == ["P001", "P002", "P003"]  # P003 corrigé


def test_demo_ignores_the_journal_and_leaves_it_unchanged(demo_project):
    journal = demo_project / "anomalies" / "corrections.csv"
    before = journal.read_bytes()

    assert main.main(["--auto", "--demo"]) == 0

    assert payment_ids(demo_project) == ["P001", "P002"]  # erreur non corrigée
    assert journal.read_bytes() == before


def test_demo_opens_the_window_without_journal(demo_project, monkeypatch):
    calls = []

    def fake_window(use_journal=True):
        calls.append(use_journal)
        validation.auto_validate(use_journal=use_journal)
        return True

    monkeypatch.setattr(main.validation_gui, "gui_validate", fake_window)
    assert main.main(["--demo"]) == 0
    assert calls == [False]


@pytest.mark.parametrize("extra", [["--source", "exports"], ["--skip-import"]])
def test_demo_is_incompatible_with_other_sources(extra):
    with pytest.raises(SystemExit):
        main.main(["--demo", *extra])
